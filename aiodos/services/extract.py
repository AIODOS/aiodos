"""Extraction de texte depuis tout type de fichier (repris et étendu de la v1).

classify() range un fichier dans une famille ; extract_text() en tire le texte.
Images et audio/vidéo passent par services.media (vision / Whisper), seulement
si l'appelant l'autorise : ce sont des traitements lents.
readable_documents() dit quels documents cette machine sait lire : un nœud
minimal ou un téléphone peut manquer d'un paquet de lecture (python-pptx…).
Les documents qui sont des archives zip (docx, pptx, xlsx, odt, epub) ne sont ouverts
que si ce qu'ils annoncent une fois décompressés reste borné (services/archives.py) ;
why_unread() dit alors pourquoi un document n'a rien donné.

HTML et RTF sont analysés en entier, en mémoire (relecture finale du 21/10/2026) : au-delà
de MAX_MARKUP_BYTES (la borne d'une partie XML), ils ne sont pas lus. Mesuré sur 8 Mo :
html.parser monte à 18 fois la taille du fichier, striprtf à 80 fois sur des groupes « { »
imbriqués (il empile un état par groupe ouvert) ; un RTF imbriqué sur plus de
MAX_RTF_DEPTH niveaux n'est pas lu non plus, ce que dit un parcours linéaire. Et le texte
d'un .odt se tire par une expression en temps linéaire : « <[^>]+> » revenait en arrière
sur une suite de « < » sans « > » (une journée de calcul pour 16 Mo).
"""
import codecs
import functools
import importlib.util
import re
import time
import zipfile
from pathlib import Path

from aiodos.i18n import t
from aiodos.services import archives

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".py", ".json", ".js", ".mjs", ".ts",
    ".tsx", ".jsx", ".css", ".scss", ".csv", ".tsv", ".log", ".xml", ".yaml",
    ".yml", ".ini", ".cfg", ".toml", ".c", ".cpp", ".h", ".hpp", ".java",
    ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".env",
    ".tex", ".srt", ".vtt", ".lua", ".kt", ".swift", ".r", ".m", ".vue",
}
DOC_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".rtf", ".epub",
           ".html", ".htm", ".odt"}
# Module qu'il faut pour lire chaque document (cf. _extract_document). HTML, EPUB
# (services/epub.py) et OpenDocument se lisent avec Python seul, comme les fichiers texte.
READER_MODULES = {".pdf": "pypdf", ".docx": "docx", ".pptx": "pptx",
                  ".xlsx": "openpyxl", ".xlsm": "openpyxl", ".rtf": "striprtf"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".tif",
             ".webp", ".heic"}
MEDIA_EXT = {".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a", ".opus",
             ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv", ".flv"}
BINARY_EXT = {".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz", ".exe",
              ".dll", ".msi", ".bin", ".so", ".dylib", ".db", ".sqlite",
              ".psd", ".ai", ".eps", ".class", ".pyc", ".pyd", ".iso", ".ico",
              ".ttf", ".otf", ".woff", ".woff2", ".blend", ".fbx", ".obj",
              ".drp", ".dra", ".lnk"}

MAX_TEXT_BYTES = 100 * 1024 * 1024     # au-delà, un document n'est pas lu
MARKUP_EXT = (".html", ".htm", ".rtf")  # analysés en entier, en mémoire
MAX_MARKUP_BYTES = archives.MAX_XML_PART
MAX_RTF_DEPTH = 1000                   # groupes { } imbriqués d'un RTF, au plus
_RTF_BRACES = re.compile(r"\\[\\{}]|[{}]")      # accolades, et leurs formes échappées
_TAG = re.compile(r"<[^<>]+>")         # une balise ; en temps linéaire (sans retour arrière)


def classify(path) -> str:
    ext = Path(path).suffix.lower()
    if ext in TEXT_EXT:
        return "text"
    if ext in DOC_EXT:
        return "document"
    if ext in IMAGE_EXT:
        return "image"
    if ext in MEDIA_EXT:
        return "media"
    if ext in BINARY_EXT:
        return "binary"
    return "unknown"


def readable_documents() -> list:
    """Formats de document (extensions) que cette machine sait lire : ceux dont
    le paquet de lecture est installé, et ceux qui n'en demandent pas."""
    return sorted(ext for ext in DOC_EXT
                  if ext not in READER_MODULES or _installed(READER_MODULES[ext]))


@functools.cache
def _installed(module) -> bool:
    """Le paquet est-il installé ? Retenu pour la vie du processus : AIODOS n'installe pas
    de paquet en route (une mise à jour le relance), et la question revient pour chaque
    document qu'un ami pourrait calculer (network/work.lease : 4,8 ms à chaque fois,
    jusqu'à 2 000 unités par demande ; relecture finale du 21/10/2026)."""
    try:
        return importlib.util.find_spec(module) is not None
    except (ImportError, ValueError):        # module rendu introuvable, ou sans spec
        return False


def paced(items, pace=None):
    """Les éléments de `items`, avec pace(secondes) après chacun, le temps qu'il a pris
    pour être traité : la cadence du travail de fond (jobs/pacing.py) passe entre deux
    pages, deux diapositives, deux chapitres."""
    began = time.monotonic()
    for item in items:
        yield item
        if pace is not None:
            pace(time.monotonic() - began)
            began = time.monotonic()


def extract_text(path, vision=None, transcribe=None, pace=None):
    """Texte du fichier, ou None s'il n'y a rien d'exploitable.

    vision     : callable(path) -> str pour les images (None = ignorer)
    transcribe : callable(path) -> str pour l'audio/vidéo (None = ignorer)
    pace       : pace(secondes) entre deux pages (PDF), diapositives, feuilles ou
                 chapitres (EPUB) d'un document (cf. paced)
    """
    p = Path(path)
    family = classify(p)
    if family == "image":
        return vision(p) if vision else None
    if family == "media":
        return transcribe(p) if transcribe else None
    if family == "binary":
        return None
    if p.stat().st_size > MAX_TEXT_BYTES:
        return None
    if family == "document":
        return _extract_document(p, pace)
    if family == "text":
        return p.read_text(encoding="utf-8", errors="ignore")
    return _text_fallback(p)


def why_unread(path) -> str | None:
    """Pourquoi extract_text n'a rien rendu, quand c'est une limite d'AIODOS et non le
    fichier qui est vide : un document trop gros, ou une archive refusée (sa raison). Dite
    à l'utilisateur avec l'unité sautée, plutôt qu'un simple « aucun texte exploitable »."""
    p = Path(path)
    if classify(p) in ("image", "media", "binary"):
        return None
    try:
        if p.stat().st_size > MAX_TEXT_BYTES:
            return t("extract.file_too_big", size=MAX_TEXT_BYTES // (1024 * 1024))
    except OSError:
        return None
    if p.suffix.lower() in archives.ZIP_DOCUMENTS:
        reason = archives.refusal(p)
        if reason is not None:
            return t("extract.archive_refused", reason=reason)
    if p.suffix.lower() in MARKUP_EXT:
        try:
            text = _markup_text(p)
        except OSError:
            return None
        if text is None:
            return t("extract.markup_too_big", size=MAX_MARKUP_BYTES // (1024 * 1024))
        if p.suffix.lower() == ".rtf" and _rtf_too_deep(text):
            return t("extract.rtf_too_deep", depth=MAX_RTF_DEPTH)
    return None


def _markup_text(p: Path) -> str | None:
    """Le texte d'un HTML ou d'un RTF (comme Path.read_text), s'il ne dépasse pas
    MAX_MARKUP_BYTES ; None sinon. Relu borné : le fichier a pu grandir depuis."""
    if p.stat().st_size > MAX_MARKUP_BYTES:
        return None
    with open(p, encoding="utf-8", errors="ignore") as f:
        text = f.read(MAX_MARKUP_BYTES + 1)
    return None if len(text) > MAX_MARKUP_BYTES else text


def _rtf_too_deep(text: str) -> bool:
    """Le RTF imbrique-t-il ses groupes sur plus de MAX_RTF_DEPTH niveaux ? En un seul
    passage, sur les seules accolades (une accolade échappée, « \\{ », n'en est pas une)."""
    depth = 0
    for match in _RTF_BRACES.finditer(text):
        brace = match.group()
        if brace == "{":
            depth += 1
            if depth > MAX_RTF_DEPTH:
                return True
        elif brace == "}" and depth:
            depth -= 1
    return False


def _extract_document(p: Path, pace=None):
    ext = p.suffix.lower()
    if ext in archives.ZIP_DOCUMENTS and not archives.readable(p):
        return None                  # archive piégée ou abîmée : rien d'exploitable
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(p))
        return "\n".join((page.extract_text() or "") for page in paced(reader.pages, pace))
    if ext == ".docx":
        import docx
        doc = docx.Document(str(p))
        parts = [par.text for par in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    # Les en-têtes de diapositive et de feuille font partie du texte extrait : le même sur
    # chaque machine, quelle que soit sa langue (un ami rend exactement ce découpage,
    # index.sample_embedded)
    if ext == ".pptx":
        from pptx import Presentation
        parts = []
        for number, slide in enumerate(paced(Presentation(str(p)).slides, pace), 1):
            parts.append(f"# Diapositive {number}")         # i18n: tel quel (texte extrait)
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts = []
        for ws in paced(wb.worksheets, pace):
            parts.append(f"# Feuille : {ws.title}")          # i18n: tel quel (texte extrait)
            # Les dimensions annoncées par la feuille ne comptent pas : openpyxl comblerait
            # jusqu'à elles chaque ligne et les lignes manquantes (A1:XFD1048576 dans un
            # fichier de quelques Ko : des milliards de cases vides parcourues)
            ws.reset_dimensions()
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext == ".rtf":
        from striprtf.striprtf import rtf_to_text
        text = _markup_text(p)
        if text is None or _rtf_too_deep(text):
            return None               # la raison : why_unread
        return rtf_to_text(text)
    if ext == ".epub":
        from aiodos.services import epub
        return epub.read_text(p, pace)
    if ext in (".html", ".htm"):
        from aiodos.services.web import html_to_text
        text = _markup_text(p)
        return None if text is None else html_to_text(text)[1]
    if ext == ".odt":
        with zipfile.ZipFile(str(p)) as z:
            try:
                with z.open("content.xml") as f:            # lecture bornée, en plus du
                    raw = f.read(archives.MAX_XML_PART + 1)  # contrôle de l'archive
            except KeyError:
                return None
        if len(raw) > archives.MAX_XML_PART:
            return None
        return _TAG.sub(" ", raw.decode("utf-8", "ignore"))
    return None


def _mostly_printable(text: str) -> bool:
    sample = text[:4000]
    printable = sum(c.isprintable() or c in "\n\r\t " for c in sample)
    return bool(sample) and printable / len(sample) > 0.85


def looks_like_text(path) -> bool:
    """Une extension inconnue est-elle du texte, à en juger par ses 4 premiers Ko (la règle
    de _text_fallback) ? Pour écarter un binaire (wallet.dat…) avant d'en faire une unité :
    il ne donnerait aucun texte, et partirait sinon tel quel chez la machine qui calcule."""
    try:
        with open(path, "rb") as f:
            head = f.read(4096)
    except OSError:
        return False
    if b"\x00" in head:
        return False
    text = codecs.getincrementaldecoder("utf-8")("ignore").decode(head, final=False)
    return _mostly_printable(text) or _mostly_printable(head.decode("latin-1"))


def _text_fallback(p: Path):
    """Extension inconnue : on la lit comme du texte si c'en est vraiment."""
    try:
        raw = p.read_bytes()
    except OSError:
        return None
    if b"\x00" in raw[:4096]:                   # octet nul : fichier binaire
        return None
    for encoding in ("utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if _mostly_printable(text):
            return text
    return None
