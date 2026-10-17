"""Extraction de fichiers (tous formats) et rendu Markdown sûr."""
import importlib.util

from aiodos import render
from aiodos.services import extract


def test_classify():
    assert extract.classify("a.py") == "text"
    assert extract.classify("a.PDF") == "document"
    assert extract.classify("a.jpg") == "image"
    assert extract.classify("a.mp4") == "media"
    assert extract.classify("a.zip") == "binary"
    assert extract.classify("a.weird") == "unknown"


def test_office_documents(tmp_path):
    import docx
    import openpyxl
    from pptx import Presentation
    doc = docx.Document()
    doc.add_paragraph("Rapport annuel 2024")
    doc.save(tmp_path / "r.docx")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Plan marketing"
    prs.save(tmp_path / "p.pptx")
    wb = openpyxl.Workbook()
    wb.active.append(["Produit", "Prix"])
    wb.save(tmp_path / "d.xlsx")
    assert "Rapport annuel 2024" in extract.extract_text(tmp_path / "r.docx")
    assert "Plan marketing" in extract.extract_text(tmp_path / "p.pptx")
    assert "Produit | Prix" in extract.extract_text(tmp_path / "d.xlsx")


def test_text_like_and_binary(tmp_path):
    (tmp_path / "page.html").write_text("<html><head><title>T</title></head>"
                                        "<body><p>Bonjour</p><script>x()</script></body></html>",
                                        encoding="utf-8")
    (tmp_path / "conf.weird").write_text("cle=valeur", encoding="utf-8")
    (tmp_path / "blob.weird").write_bytes(bytes(range(256)))
    (tmp_path / "lettre.rtf").write_text(r"{\rtf1\ansi Bonjour RTF.}", encoding="utf-8")
    assert extract.extract_text(tmp_path / "page.html").strip() == "Bonjour"
    assert extract.extract_text(tmp_path / "conf.weird") == "cle=valeur"
    assert extract.extract_text(tmp_path / "blob.weird") is None
    assert "Bonjour RTF" in extract.extract_text(tmp_path / "lettre.rtf")


def test_readable_documents_follow_the_installed_packages(monkeypatch):
    """Un nœud minimal ou un téléphone peut manquer d'un paquet de lecture : il ne
    lit alors pas ce format. HTML et OpenDocument se lisent toujours."""
    assert extract.readable_documents() == sorted(extract.DOC_EXT)    # requirements.txt
    find_spec = importlib.util.find_spec
    monkeypatch.setattr(importlib.util, "find_spec", lambda name, package=None:
                        None if name == "pptx" else find_spec(name, package))
    assert extract.readable_documents() == sorted(extract.DOC_EXT - {".pptx"})
    monkeypatch.setattr(importlib.util, "find_spec", lambda name, package=None: None)
    assert extract.readable_documents() == [".epub", ".htm", ".html", ".odt"]


def test_media_needs_explicit_handlers(tmp_path):
    image = tmp_path / "photo.png"
    image.write_bytes(b"\x89PNG fake")
    assert extract.extract_text(image) is None
    assert extract.extract_text(image, vision=lambda p: "texte lu") == "texte lu"


def test_render_escapes_raw_html_and_unsafe_links():
    html = render.to_html('<script>alert(1)</script>\n\n[x](javascript:alert(1))')
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert 'href="javascript' not in html


def test_render_code_and_links():
    html = render.to_html("```python\nprint('ok')\n```\n\n[site](https://example.org)")
    assert 'class="code-block"' in html and "Copier" in html and "python" in html
    assert 'target="_blank"' in html and 'rel="noopener noreferrer"' in html
