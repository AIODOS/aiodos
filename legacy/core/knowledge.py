"""Base de connaissance locale (RAG), organisée en bibliothèques (collections).

- Chaque passage ("chunk") appartient à une **collection** nommée.
- Une collection peut être **activée/désactivée**.
- On peut indexer un fichier, plusieurs fichiers, un **dossier entier**, ou des URLs.

⚠ Concurrence : plusieurs threads peuvent écrire (import + indexation des
conversations). Toute lecture-modification-écriture du store passe par _LOCK,
et on **relit** le fichier juste avant d'écrire pour ne rien écraser.

Stockage : data/knowledge/chunks.json + data/knowledge/collections.json
Recherche : similarité cosinus (numpy).
"""
import json
import hashlib
import threading
from pathlib import Path

import numpy as np

from core import config, ollama_client

KNOW_DIR = config.DATA_DIR / "knowledge"
STORE_PATH = KNOW_DIR / "chunks.json"
COLL_PATH = KNOW_DIR / "collections.json"

DEFAULT_COLLECTION = "Général"
CONV_COLLECTION = "💬 Conversations"   # collection spéciale : contenu des conversations

_LOCK = threading.RLock()              # sérialise tous les accès en écriture au store

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".py", ".json", ".js", ".ts", ".tsx",
    ".jsx", ".html", ".htm", ".css", ".csv", ".log", ".xml", ".yaml", ".yml",
    ".ini", ".cfg", ".toml", ".c", ".cpp", ".h", ".hpp", ".java", ".cs", ".go",
    ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".env", ".tex",
}

# Documents nécessitant une extraction de texte dédiée
DOC_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".rtf", ".epub",
           ".html", ".htm", ".odt"}
SUPPORTED_EXT = TEXT_EXT | DOC_EXT

# Fichiers sans texte exploitable directement (sautés sauf OCR/transcription)
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".tif",
             ".webp", ".svg", ".ico", ".heic"}
MEDIA_EXT = {".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a", ".opus",
             ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv", ".flv"}
BINARY_EXT = {".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".exe", ".dll",
              ".bin", ".so", ".dylib", ".db", ".sqlite", ".psd", ".ai",
              ".eps", ".class", ".pyc"}
NO_TEXT_EXT = IMAGE_EXT | MEDIA_EXT | BINARY_EXT

CHUNK_SIZE = 900
CHUNK_OVERLAP = 150


# --------------------------------------------------------------------- stockage
def _load():
    if STORE_PATH.exists():
        try:
            return json.loads(STORE_PATH.read_text(encoding="utf-8"))
        except Exception:
            return []
    return []


def _save(chunks):
    KNOW_DIR.mkdir(parents=True, exist_ok=True)
    STORE_PATH.write_text(json.dumps(chunks, ensure_ascii=False), encoding="utf-8")


def _load_collections():
    data = {}
    if COLL_PATH.exists():
        try:
            data = json.loads(COLL_PATH.read_text(encoding="utf-8"))
        except Exception:
            data = {}
    if DEFAULT_COLLECTION not in data:
        data[DEFAULT_COLLECTION] = {"enabled": True}
    return data


def _save_collections(data):
    KNOW_DIR.mkdir(parents=True, exist_ok=True)
    COLL_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                         encoding="utf-8")


def _coll_of(chunk):
    return chunk.get("collection", DEFAULT_COLLECTION)


def _id(collection, source, idx):
    return hashlib.md5(f"{collection}:{source}:{idx}".encode()).hexdigest()


def _chunk_text(text):
    text = text.strip()
    out, i = [], 0
    step = max(1, CHUNK_SIZE - CHUNK_OVERLAP)
    while i < len(text):
        piece = text[i:i + CHUNK_SIZE]
        if piece.strip():
            out.append(piece)
        i += step
    return out


def _text_fallback(p, ext):
    """Dernier recours : lire un fichier inconnu comme du texte, si c'est du texte."""
    if ext in NO_TEXT_EXT:
        return None
    try:
        if p.stat().st_size > 25 * 1024 * 1024:    # > 25 Mo : on évite
            return None
        raw = p.read_bytes()
    except Exception:
        return None
    if b"\x00" in raw[:4096]:                       # octet nul = binaire
        return None
    for enc in ("utf-8", "latin-1"):
        try:
            txt = raw.decode(enc)
        except Exception:
            continue
        sample = txt[:4000]
        if sample and sum(c.isprintable() or c in "\n\r\t " for c in sample) \
                / len(sample) > 0.85:
            return txt
    return None


def _extract_text(path):
    """Extrait le texte de TOUT type de fichier. None si aucun texte exploitable."""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            return "\n".join((pg.extract_text() or "") for pg in reader.pages)
        if ext == ".docx":
            import docx
            doc = docx.Document(str(p))
            parts = [par.text for par in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(p))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            return "\n".join(parts)
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"# Feuille : {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        if ext == ".rtf":
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(p.read_text(encoding="utf-8", errors="ignore"))
        if ext == ".epub":
            import ebooklib
            from ebooklib import epub
            from core import web
            book = epub.read_epub(str(p))
            parts = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                ex = web._TextExtractor()
                ex.feed(item.get_content().decode("utf-8", "ignore"))
                parts.append(ex.text())
            return "\n".join(parts)
        if ext in (".html", ".htm"):
            from core import web
            ex = web._TextExtractor()
            ex.feed(p.read_text(encoding="utf-8", errors="ignore"))
            return ex.text()
        if ext == ".odt":
            import zipfile
            import re as _re
            with zipfile.ZipFile(str(p)) as z:
                xml = z.read("content.xml").decode("utf-8", "ignore")
            return _re.sub(r"<[^>]+>", " ", xml)
        if ext in TEXT_EXT:
            return p.read_text(encoding="utf-8", errors="ignore")
        if ext in IMAGE_EXT:
            if config.get("media_vision"):
                from core import media
                txt = media.vision_extract(p, mode=config.get("media_vision_mode"))
                return f"[Image : {p.name}]\n{txt}" if txt else None
            return None
        if ext in MEDIA_EXT:
            if config.get("media_transcribe"):
                from core import media
                return media.transcribe(p)
            return None
    except Exception:
        return None
    return _text_fallback(p, ext)


# ----------------------------------------------------------------- collections
def create_collection(name):
    name = name.strip()
    if not name:
        return
    with _LOCK:
        colls = _load_collections()
        if name not in colls:
            colls[name] = {"enabled": True}
            _save_collections(colls)


def delete_collection(name):
    with _LOCK:
        _save([c for c in _load() if _coll_of(c) != name])
        colls = _load_collections()
        colls.pop(name, None)
        if name == DEFAULT_COLLECTION:
            colls[DEFAULT_COLLECTION] = {"enabled": True}
        _save_collections(colls)


def set_collection_enabled(name, enabled):
    with _LOCK:
        colls = _load_collections()
        colls.setdefault(name, {})["enabled"] = bool(enabled)
        _save_collections(colls)


def _enabled_set():
    return {n for n, m in _load_collections().items() if m.get("enabled", True)}


def list_collections(include_conv=False):
    """{nom: {enabled, chunks, sources}}. Exclut Conversations par défaut."""
    colls = _load_collections()
    info = {n: {"enabled": m.get("enabled", True), "chunks": 0, "sources": set()}
            for n, m in colls.items() if n != CONV_COLLECTION or include_conv}
    for c in _load():
        name = _coll_of(c)
        if name == CONV_COLLECTION and not include_conv:
            continue
        info.setdefault(name, {"enabled": True, "chunks": 0, "sources": set()})
        info[name]["chunks"] += 1
        info[name]["sources"].add(c["source"])
    for v in info.values():
        v["sources"] = sorted(v["sources"])
    return info


def list_entries(include_conv=False):
    """[{collection, source, kind, chunks}] — une ligne par source."""
    agg = {}
    for c in _load():
        coll = _coll_of(c)
        if coll == CONV_COLLECTION and not include_conv:
            continue
        key = (coll, c["source"])
        e = agg.setdefault(key, {"collection": coll, "source": c["source"],
                                 "kind": c.get("kind", "file"), "chunks": 0})
        e["chunks"] += 1
    return sorted(agg.values(), key=lambda e: (e["collection"], e["source"]))


def rename_collection(old, new):
    new = new.strip()
    if not new or old == new or old == CONV_COLLECTION:
        return
    with _LOCK:
        chunks = _load()
        for c in chunks:
            if _coll_of(c) == old:
                c["collection"] = new
        _save(chunks)
        colls = _load_collections()
        meta = colls.pop(old, {"enabled": True})
        colls[new] = meta
        if old == DEFAULT_COLLECTION:
            colls[DEFAULT_COLLECTION] = {"enabled": True}
        _save_collections(colls)


def count():
    return len(_load())


def manifest(libraries=True, conversations=False, per_coll=25):
    """Sommaire des sources disponibles, pour informer l'IA de ce qu'elle possède."""
    enabled = _enabled_set()
    by_coll = {}
    for e in list_entries(include_conv=conversations):
        coll = e["collection"]
        if coll == CONV_COLLECTION:
            if not conversations:
                continue
        elif not libraries or coll not in enabled:
            continue
        by_coll.setdefault(coll, []).append(e["source"])
    lines = []
    for coll, srcs in by_coll.items():
        shown = ", ".join(srcs[:per_coll])
        more = f" … (+{len(srcs) - per_coll})" if len(srcs) > per_coll else ""
        lines.append(f"• {coll} ({len(srcs)} éléments) : {shown}{more}")
    return "\n".join(lines)


# ------------------------------------------------------------------ indexation
def add_text(text, source, collection=DEFAULT_COLLECTION, kind="file", progress=None):
    """Indexe un texte. Remplace l'ancienne version (même source+collection)."""
    create_collection(collection)
    pieces = _chunk_text(text)
    new = []
    for idx, piece in enumerate(pieces):
        new.append({
            "id": _id(collection, source, idx),
            "source": source, "collection": collection, "kind": kind,
            "text": piece, "embedding": ollama_client.embed(piece),
        })
        if progress:
            progress(idx + 1, len(pieces))
    with _LOCK:                                    # relecture + écriture atomique
        chunks = [c for c in _load()
                  if not (c["source"] == source and _coll_of(c) == collection)]
        chunks.extend(new)
        _save(chunks)
    return len(pieces)


def add_file(path, collection=DEFAULT_COLLECTION, progress=None):
    p = Path(path)
    text = _extract_text(p)             # tout type accepté
    if not text or not text.strip():
        raise ValueError(
            "Aucun texte exploitable (image, audio/vidéo, ou fichier binaire). "
            "Ce type nécessite l'OCR ou la transcription.")
    return add_text(text, p.name, collection, kind="file", progress=progress)


def iter_text_files(folder, recursive=True):
    folder = Path(folder)
    it = folder.rglob("*") if recursive else folder.glob("*")
    for p in sorted(it):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
            yield p


def _index_files(files, collection, root=None, kind="file", file_progress=None):
    """Indexe une liste de fichiers. Les embeddings (lents) sont calculés HORS
    verrou ; seule la fusion finale est verrouillée et relit le store."""
    create_collection(collection)
    total = len(files)
    added = 0
    new = []
    sources_done = set()
    for fi, fpath in enumerate(files):
        p = Path(fpath)
        try:
            text = _extract_text(p)
        except Exception:
            text = None
        if not text or not text.strip():
            continue
        if root is not None:
            try:
                source = str(p.relative_to(root))
            except ValueError:
                source = p.name
        else:
            source = p.name
        for idx, piece in enumerate(_chunk_text(text)):
            new.append({
                "id": _id(collection, source, idx),
                "source": source, "collection": collection, "kind": kind,
                "text": piece, "embedding": ollama_client.embed(piece),
            })
        sources_done.add(source)
        added += 1
        if file_progress:
            file_progress(fi + 1, total, source)
    with _LOCK:
        chunks = [c for c in _load()
                  if not (_coll_of(c) == collection and c["source"] in sources_done)]
        chunks.extend(new)
        _save(chunks)
    return added, total


def add_files(paths, collection=DEFAULT_COLLECTION, file_progress=None):
    return _index_files(list(paths), collection, root=None,
                        kind="file", file_progress=file_progress)


def add_folder(folder, collection=DEFAULT_COLLECTION, recursive=True, file_progress=None):
    folder = Path(folder)
    allf = sorted(p for p in (folder.rglob("*") if recursive else folder.glob("*"))
                  if p.is_file())
    if not allf:
        raise ValueError("Dossier vide.")
    # on tente TOUS les fichiers ; ceux sans texte exploitable sont ignorés
    added, total = _index_files(allf, collection, root=folder,
                                kind="folder", file_progress=file_progress)
    if added == 0:
        from collections import Counter
        exts = Counter((p.suffix.lower() or "(sans extension)") for p in allf)
        found = ", ".join(f"{e} ({n})" for e, n in exts.most_common(10))
        raise ValueError(
            "Aucun texte exploitable dans ce dossier.\n\n"
            f"Types trouvés : {found}\n\n"
            "Documents et code sont indexés ; les images, audio et vidéo "
            "nécessitent l'OCR ou la transcription (pas encore actifs).")
    return added, total


def add_urls(urls, collection=DEFAULT_COLLECTION, file_progress=None):
    """Télécharge des pages web et les indexe (kind='url'). source = URL."""
    from core import web
    create_collection(collection)
    urls = [u.strip() for u in urls if u.strip()]
    total = len(urls)
    added = 0
    for ui, url in enumerate(urls):
        try:
            title, text = web.fetch_url_text(url)
        except Exception:
            if file_progress:
                file_progress(ui + 1, total, f"(échec) {url}")
            continue
        if not text.strip():
            continue
        add_text(f"{title}\n\n{text}", url, collection, kind="url")
        added += 1
        if file_progress:
            file_progress(ui + 1, total, url)
    return added, total


def remove_source(source, collection=None):
    with _LOCK:
        if collection is None:
            _save([c for c in _load() if c["source"] != source])
        else:
            _save([c for c in _load()
                   if not (c["source"] == source and _coll_of(c) == collection)])


def clear():
    with _LOCK:
        _save([])


# ----------------------------------------------------- index des conversations
def index_conversation(sid, title, messages):
    """Indexe (incrémental) une conversation dans CONV_COLLECTION, sans écraser
    le reste du store (relecture sous verrou)."""
    create_collection(CONV_COLLECTION)
    with _LOCK:
        existing = {(c.get("sid"), c.get("mi")) for c in _load()
                    if _coll_of(c) == CONV_COLLECTION}
    new = []
    for mi, m in enumerate(messages):
        if (sid, mi) in existing:
            continue
        who = "Utilisateur" if m["role"] == "user" else "Assistant"
        text = f"{who} : {m['content']}"
        for pidx, piece in enumerate(_chunk_text(text) or [text]):
            new.append({
                "id": _id(CONV_COLLECTION, sid, f"{mi}.{pidx}"),
                "source": title, "collection": CONV_COLLECTION,
                "kind": "conversation", "sid": sid, "mi": mi,
                "text": piece, "embedding": ollama_client.embed(piece),
            })
    with _LOCK:
        chunks = _load()
        changed = False
        for c in chunks:                       # rafraîchir le titre affiché
            if _coll_of(c) == CONV_COLLECTION and c.get("sid") == sid \
                    and c.get("source") != title:
                c["source"] = title
                changed = True
        present = {(c.get("sid"), c.get("mi")) for c in chunks
                   if _coll_of(c) == CONV_COLLECTION}
        for c in new:
            if (c["sid"], c["mi"]) not in present:
                chunks.append(c)
                present.add((c["sid"], c["mi"]))
                changed = True
        if changed:
            _save(chunks)


# -------------------------------------------------------------------- recherche
def search(query, k=4, libraries=True, conversations=False):
    """k passages les plus proches selon les sources autorisées."""
    enabled = _enabled_set()

    def allowed(coll):
        if coll == CONV_COLLECTION:
            return conversations
        return libraries and coll in enabled

    chunks = [c for c in _load() if allowed(_coll_of(c))]
    if not chunks:
        return []
    q = np.array(ollama_client.embed(query), dtype=np.float32)
    qn = float(np.linalg.norm(q)) or 1.0
    scored = []
    for c in chunks:
        v = np.array(c["embedding"], dtype=np.float32)
        vn = float(np.linalg.norm(v)) or 1.0
        sim = float(np.dot(q, v) / (qn * vn))
        scored.append((sim, c))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [{"source": c["source"], "collection": _coll_of(c), "text": c["text"],
             "score": s} for s, c in scored[:k]]
